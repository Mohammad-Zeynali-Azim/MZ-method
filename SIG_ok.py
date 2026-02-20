#Similarity-Driven Division by Two Graph Complex Networ or SiCNet---With the Saving in Excel File
import MG9846 as mg
import numpy as np
import pylab as plt
import networkx as nx
import random
import collections
import pandas as pd
from docx import Document
from io import BytesIO
from matplotlib.backends.backend_agg import FigureCanvasAgg as FigureCanvas
from sklearn.mixture import GaussianMixture
from statsmodels.distributions.empirical_distribution import ECDF
from scipy.optimize import curve_fit
import community  # Import the community library
import matplotlib.cm as cm  # Import the colormap module

# Create a Word file
doc = Document()
doc.add_heading('Results of Proposed Model', 0)

def generate_random_numbers():
    n_digits = random.randint(2, 6)  # n between 10 and 999999
    n = random.randint(10**(n_digits-1), 10**n_digits - 1)
    m_digits = n_digits + 4
    m = random.randint(10**(m_digits-1), 10**m_digits - 1)
    return n, m

def create_graph(n, m):
    num1 = [int(d) for d in str(n)]
    L1 = mg.Graph_Generate(num1, 200)
    mg.After_constructing(L1)

    natije1 = []
    for j in range(1, 2000):
        L3 = mg.Dimond_value(j)
        if any(x != 0 for x in L3):
            natije1.append(L3)

    num2 = [int(d) for d in str(m)]
    L2 = mg.Graph_Generate(num2, 200)
    mg.After_constructing(L2)

    natije2 = []
    for j in range(1, 2000):
        L4 = mg.Dimond_value(j)
        natije2.append(L4)
    list_for_Cyto1=[]
    list_for_Cyto2=[]
    G = nx.Graph()
    for k in range(len(natije1)):
        for m in range(len(natije2)):
            if natije1[k] == natije2[m]:
                G.add_edge(k, m)
                list_for_Cyto1.append(k)
                list_for_Cyto2.append(m)
    return G, list_for_Cyto1, list_for_Cyto2

#n, m = generate_random_numbers()
#print(n,m)
G, list_for_Cyto1, list_for_Cyto2 = create_graph(6543, 12345678)
print(G.edges)
import xlsxwriter
i=1;
j=1;
workbook = xlsxwriter.Workbook('u45576776_12345678.xlsx')
worksheet = workbook.add_worksheet()
for item in list_for_Cyto1:
    worksheet.write(i, 0, item)     # Writes an int
    i=i+1;
for item1 in list_for_Cyto2:
    worksheet.write(j, 1, item1)     # Writes an int
    j=j+1;

workbook.close()
print("hhhhhhhhhhhhh")